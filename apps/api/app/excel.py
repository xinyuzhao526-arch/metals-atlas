from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


SHEETS: dict[str, list[str]] = {
    "projects": ["record_id", "row_version", "slug", "name", "country_iso3", "operator_company", "latitude", "longitude", "status", "raw_material_route"],
    "companies": ["record_id", "row_version", "canonical_name", "legal_name", "country_iso3", "website", "fiscal_year_start_month"],
    "ownership": ["record_id", "row_version", "project_slug", "company_name", "ownership_pct", "valid_from", "valid_to"],
    "production": ["record_id", "row_version", "project_slug", "metal_code", "source_code", "record_key", "original_value", "original_unit", "normalized_value", "normalized_unit", "missing_reason", "period_start", "period_end", "effective_date", "calendar_basis", "fiscal_year_label", "fiscal_year_start_month", "period_type", "ownership_basis", "production_stage", "source_published_at", "source_locator", "verified_at", "is_cumulative", "is_estimate", "notes"],
    "guidance": ["record_id", "row_version", "project_slug", "metal_code", "source_code", "record_key", "guidance_low", "guidance_high", "original_unit", "normalized_unit", "missing_reason", "period_start", "period_end", "effective_date", "calendar_basis", "fiscal_year_label", "fiscal_year_start_month", "period_type", "ownership_basis", "production_stage", "source_published_at", "source_locator", "verified_at", "guidance_kind", "adjustment_type", "adjustment_date", "notes"],
    "reserves": ["record_id", "row_version", "project_slug", "metal_code", "source_code", "record_key", "contained_metal_value", "original_unit", "normalized_unit", "missing_reason", "period_start", "period_end", "effective_date", "calendar_basis", "fiscal_year_label", "fiscal_year_start_month", "period_type", "ownership_basis", "production_stage", "source_published_at", "source_locator", "verified_at", "reserve_kind", "classification", "ore_tonnage", "grade_pct", "notes"],
    "sources": ["record_id", "row_version", "code", "organization_name", "material_title", "material_url", "published_at", "source_type", "verified_at", "is_demo", "notes"],
}

STATIC_LOOKUPS = {
    "metal_code": ["Cu", "Al", "Pb", "Zn", "Ni", "Sn", "Li"],
    "calendar_basis": ["calendar_year", "fiscal_year"],
    "period_type": ["quarter", "ytd", "annual", "other"],
    "ownership_basis": ["project_100", "equity", "attributable", "consolidated", "unknown"],
    "production_stage": ["mine_contained_metal", "concentrate_contained_metal", "cathode", "anode_blister", "smelter_output"],
    "source_type": ["quarterly_report", "annual_report", "announcement", "data_webpage", "demo_fixture"],
}


def lookup_values(country_iso3s: Iterable[str]) -> dict[str, list[str]]:
    countries = sorted(
        {str(code).strip().upper() for code in country_iso3s if str(code).strip()}
    )
    return {"country_iso3": countries, **STATIC_LOOKUPS}

FIELD_HELP: dict[str, tuple[str, str, str, str]] = {
    "record_id": ("数据库 UUID；新增时留空，更新时从导出文件带回", "text", "550e8400-e29b-41d4-a716-446655440000", "提供时必须与稳定键指向同一记录"),
    "row_version": ("乐观锁版本；更新时从导出文件带回", "integer", "1", "必须与数据库当前版本一致"),
    "slug": ("项目稳定英文短标识", "text", "escondida", "唯一，建议小写连字符"),
    "name": ("项目显示名称", "text", "Escondida", "不得为空"),
    "country_iso3": ("ISO 3166-1 三字母国家代码", "text", "CHL", "必须存在于国家字典"),
    "operator_company": ("运营方标准公司名", "text", "BHP", "填写时必须匹配 companies.canonical_name"),
    "latitude": ("项目纬度", "decimal", "-24.268", "-90 至 90"),
    "longitude": ("项目经度", "decimal", "-69.071", "-180 至 180"),
    "status": ("项目状态", "text", "operating", "使用约定状态代码"),
    "raw_material_route": ("原料路线", "text", "primary", "可留空"),
    "canonical_name": ("公司标准名称", "text", "BHP", "唯一且不得为空"),
    "legal_name": ("公司法定名称", "text", "BHP Group Limited", "可留空"),
    "website": ("公司官方网站", "url", "https://www.bhp.com/", "填写完整 URL"),
    "fiscal_year_start_month": ("财年起始月份", "integer", "7", "1 至 12；财年观察必填"),
    "project_slug": ("关联项目 slug", "text", "escondida", "必须匹配 projects.slug"),
    "company_name": ("关联公司标准名称", "text", "BHP", "必须匹配 companies.canonical_name"),
    "ownership_pct": ("持股比例，按百分数填写", "decimal", "57.5", "0 至 100；不得与同项目同公司有效期重叠"),
    "valid_from": ("持股有效起始日", "date", "2025-07-01", "YYYY-MM-DD"),
    "valid_to": ("持股有效结束日；持续有效时留空", "date", "2026-06-30", "YYYY-MM-DD 且不得早于 valid_from"),
    "metal_code": ("金属代码", "text", "Cu", "从 _LOOKUPS 选择"),
    "source_code": ("具体来源材料稳定代码", "text", "BHP-FY2026-Q4", "必须匹配 sources.code；与 record_key 组成幂等键"),
    "record_key": ("来源内业务事实稳定键", "text", "escondida-fy2026-production", "与 source_code 联合唯一"),
    "original_value": ("来源披露原始数值", "decimal", "123.456", "缺失时留空，不得填 0 代替"),
    "original_unit": ("来源披露原始单位", "text", "kt", "数值存在时建议填写"),
    "normalized_value": ("标准化数值", "decimal", "123.456", "缺失时留空并填写 missing_reason"),
    "normalized_unit": ("标准化单位", "text", "kt", "不得为空"),
    "missing_reason": ("结构化缺失原因", "text", "not_disclosed", "数值为空时必填"),
    "period_start": ("统计期间起始日", "date", "2025-07-01", "YYYY-MM-DD"),
    "period_end": ("统计期间结束日", "date", "2026-06-30", "YYYY-MM-DD 且不得早于期间起始日"),
    "effective_date": ("数据有效日期", "date", "2026-06-30", "YYYY-MM-DD"),
    "calendar_basis": ("自然年或公司财年", "enum", "fiscal_year", "从 _LOOKUPS 选择"),
    "fiscal_year_label": ("公司财年标签", "text", "FY2026", "calendar_basis=fiscal_year 时必填"),
    "period_type": ("季度、YTD、年度或其他期间类型", "enum", "annual", "从 _LOOKUPS 选择"),
    "ownership_basis": ("项目 100%、权益、应占或合并口径", "enum", "project_100", "从 _LOOKUPS 选择"),
    "production_stage": ("生产环节", "enum", "mine_contained_metal", "从 _LOOKUPS 选择；禁止跨环节合计"),
    "source_published_at": ("来源材料发布日期", "date", "2026-08-20", "YYYY-MM-DD"),
    "source_locator": ("页码、表格、段落或字段路径", "text", "p. 12 table 3", "不得为空"),
    "verified_at": ("人工核验日期", "date", "2026-09-16", "YYYY-MM-DD"),
    "is_cumulative": ("是否为累计值", "boolean", "false", "true/false"),
    "is_estimate": ("是否为估算值", "boolean", "false", "true/false"),
    "guidance_low": ("指引下限", "decimal", "120", "可为空；不得大于上限"),
    "guidance_high": ("指引上限", "decimal", "140", "可为空；不得小于下限"),
    "guidance_kind": ("初始、修订或维持指引", "enum", "initial", "initial/revision/maintained"),
    "adjustment_type": ("指引调整类型", "text", "operational", "可留空"),
    "adjustment_date": ("指引调整日期", "date", "2026-02-15", "YYYY-MM-DD，可留空"),
    "contained_metal_value": ("储量或资源量中的含金属量", "decimal", "25000", "缺失时留空并填写 missing_reason"),
    "reserve_kind": ("储量或资源量", "enum", "reserve", "reserve/resource"),
    "classification": ("储量或资源量分类", "text", "proved_and_probable", "不得为空"),
    "ore_tonnage": ("矿石量", "decimal", "100000", "可留空"),
    "grade_pct": ("品位百分比", "decimal", "0.5", "可留空"),
    "notes": ("补充说明", "text", "等待人工核验", "可留空"),
    "code": ("具体来源材料稳定代码", "text", "BHP-FY2026-Q4", "唯一且不得为空"),
    "organization_name": ("发布机构名称", "text", "BHP", "不得为空"),
    "material_title": ("具体材料标题", "text", "Operational Review FY2026", "不得为空"),
    "material_url": ("具体材料 URL", "url", "https://example.com/report.pdf", "不得使用机构主页代替具体材料"),
    "published_at": ("材料发布日期", "date", "2026-08-20", "YYYY-MM-DD"),
    "source_type": ("材料类型", "enum", "annual_report", "从 _LOOKUPS 选择"),
    "is_demo": ("是否为演示材料", "boolean", "false", "true/false"),
}

REQUIRED_FIELDS: dict[str, set[str]] = {
    "projects": {"slug", "name", "country_iso3", "status"},
    "companies": {"canonical_name", "country_iso3", "fiscal_year_start_month"},
    "ownership": {"project_slug", "company_name", "ownership_pct", "valid_from"},
    "sources": {"code", "organization_name", "material_title", "material_url", "published_at", "source_type", "verified_at", "is_demo"},
}

OBSERVATION_REQUIRED = {"project_slug", "metal_code", "source_code", "record_key", "normalized_unit", "period_start", "period_end", "effective_date", "calendar_basis", "period_type", "ownership_basis", "production_stage", "source_published_at", "source_locator", "verified_at"}


def build_template(country_iso3s: Iterable[str] = ()) -> bytes:
    lookups_data = lookup_values(country_iso3s)
    wb = Workbook()
    readme = wb.active
    readme.title = "_README"
    readme.append(["全球金属供给图谱 Phase 1A Excel 模板"])
    readme.append(["规则", "空白数值表示 NULL；missing_reason 记录原因；观察记录必须填写 source_code 与 record_key。"])
    readme.append(["流程", "上传预览 → 确认导入 → 待审核 → 接受并发布。"])
    readme.append([])
    readme.append(["工作表", "字段", "字段说明", "必填", "数据类型", "示例值", "校验规则"])
    for sheet_name, headers in SHEETS.items():
        required = REQUIRED_FIELDS.get(sheet_name, set()) | (OBSERVATION_REQUIRED if sheet_name in {"production", "guidance", "reserves"} else set())
        if sheet_name == "guidance":
            required = required | {"guidance_kind"}
        if sheet_name == "reserves":
            required = required | {"reserve_kind", "classification"}
        for field in headers:
            description, data_type, example, validation = FIELD_HELP[field]
            readme.append([sheet_name, field, description, "是" if field in required else "否/条件必填", data_type, example, validation])
    for cell in readme[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324F")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    readme.freeze_panes = "A6"
    for column, width in {"A": 15, "B": 28, "C": 48, "D": 14, "E": 14, "F": 34, "G": 52}.items():
        readme.column_dimensions[column].width = width

    lookups = wb.create_sheet("_LOOKUPS")
    for col, (name, values) in enumerate(lookups_data.items(), start=1):
        lookups.cell(1, col, name)
        for row, value in enumerate(values, start=2):
            lookups.cell(row, col, value)

    for sheet_name, headers in SHEETS.items():
        ws = wb.create_sheet(sheet_name)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17324F")
            cell.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).coordinate}"
        for index, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(1, index).column_letter].width = max(14, min(28, len(header) + 4))
        for lookup_name, lookup_items in lookups_data.items():
            if lookup_name in headers:
                col = headers.index(lookup_name) + 1
                lookup_col = list(lookups_data).index(lookup_name) + 1
                letter = lookups.cell(1, lookup_col).column_letter
                last_row = max(2, len(lookup_items) + 1)
                validation = DataValidation(type="list", formula1=f"'_LOOKUPS'!${letter}$2:${letter}${last_row}", allow_blank=True)
                ws.add_data_validation(validation)
                validation.add(f"{ws.cell(2, col).column_letter}2:{ws.cell(500, col).column_letter}500")

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def read_workbook(path: Path) -> dict[str, list[tuple[int, dict[str, Any]]]]:
    workbook = load_workbook(path, data_only=False)
    missing = [sheet for sheet in SHEETS if sheet not in workbook.sheetnames]
    if missing:
        raise ValueError(f"缺少工作表: {', '.join(missing)}")
    result: dict[str, list[tuple[int, dict[str, Any]]]] = {}
    for sheet, expected in SHEETS.items():
        ws = workbook[sheet]
        headers = [cell.value for cell in ws[1]]
        if headers != expected:
            raise ValueError(f"{sheet} 列名或顺序不正确")
        rows: list[tuple[int, dict[str, Any]]] = []
        for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if all(value is None for value in values):
                continue
            rows.append((row_number, {header: value for header, value in zip(headers, values, strict=True)}))
        result[sheet] = rows
    return result
