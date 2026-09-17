from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from app.cli import seed_countries
from app.excel import build_template
from app.imports import normalize_row
from app.models import Country


def test_fiscal_year_requires_label_and_month():
    payload, errors = normalize_row(
        "production",
        {
            "project_slug": "mine",
            "metal_code": "Cu",
            "source_code": "SRC",
            "record_key": "row",
            "normalized_value": 1,
            "normalized_unit": "kt",
            "period_start": "2025-07-01",
            "period_end": "2026-06-30",
            "effective_date": "2026-06-30",
            "calendar_basis": "fiscal_year",
            "period_type": "annual",
            "ownership_basis": "project_100",
            "production_stage": "mine_contained_metal",
            "source_published_at": "2026-07-01",
            "source_locator": "p1",
            "verified_at": "2026-09-16",
        },
        {"projects": {"mine"}, "metals": {"Cu"}, "sources": {"SRC"}},
    )
    assert any("财年记录" in error for error in errors)
    assert payload["normalized_value"] == "1"


def test_null_stays_null_and_requires_reason():
    _, errors = normalize_row(
        "production",
        {
            "project_slug": "mine",
            "metal_code": "Cu",
            "source_code": "SRC",
            "record_key": "row",
            "normalized_unit": "kt",
            "period_start": "2025-01-01",
            "period_end": "2025-12-31",
            "effective_date": "2025-12-31",
            "calendar_basis": "calendar_year",
            "period_type": "annual",
            "ownership_basis": "project_100",
            "production_stage": "mine_contained_metal",
            "source_published_at": "2026-01-01",
            "source_locator": "p1",
            "verified_at": "2026-09-16",
        },
        {"projects": {"mine"}, "metals": {"Cu"}, "sources": {"SRC"}},
    )
    assert any("missing_reason" in error for error in errors)


def project_workbook(country_iso3: str) -> bytes:
    workbook = load_workbook(BytesIO(build_template(["CHL", "USA"])))
    workbook["projects"].append(
        [None, None, "country-check", "Country Check", country_iso3, None, None, None, "unknown", None]
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def upload_project_workbook(client, headers, country_iso3: str):
    return client.post(
        "/api/v1/admin/imports",
        files={
            "file": (
                f"{country_iso3}.xlsx",
                BytesIO(project_workbook(country_iso3)),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )


def test_existing_iso3_project_row_passes_preview(authenticated):
    client, headers = authenticated
    uploaded = upload_project_workbook(client, headers, "CHL")
    assert uploaded.status_code == 200, uploaded.text
    preview = client.get(
        f"/api/v1/admin/imports/{uploaded.json()['id']}/preview"
    ).json()
    project = next(row for row in preview["rows"] if row["sheet"] == "projects")
    assert project["classification"] == "new"
    assert project["errors"] == []


def test_unknown_iso3_project_row_fails_preview(authenticated):
    client, headers = authenticated
    uploaded = upload_project_workbook(client, headers, "XXX")
    assert uploaded.status_code == 200, uploaded.text
    preview = client.get(
        f"/api/v1/admin/imports/{uploaded.json()['id']}/preview"
    ).json()
    project = next(row for row in preview["rows"] if row["sheet"] == "projects")
    assert project["sheet"] == "projects"
    assert project["row_number"] == 2
    assert project["classification"] == "invalid"
    assert project["errors"] == ["未知 ISO3: XXX"]
    assert preview["summary"] == {
        "new": 0,
        "update": 0,
        "no_change": 0,
        "conflict": 0,
        "invalid": 1,
    }


def test_downloaded_template_uses_database_country_lookup(authenticated, session):
    client, _headers = authenticated
    seed_countries(session)

    response = client.get("/api/v1/admin/excel/template")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert "countries" not in workbook.sheetnames

    lookups = workbook["_LOOKUPS"]
    headers = [cell.value for cell in lookups[1]]
    country_column = headers.index("country_iso3") + 1
    actual_codes = [
        lookups.cell(row, country_column).value
        for row in range(2, lookups.max_row + 1)
        if lookups.cell(row, country_column).value
    ]
    expected_codes = sorted(country.iso3 for country in session.scalars(select(Country)))
    assert actual_codes == expected_codes
    assert {"PER", "COD", "ZMB", "IDN", "CHN", "BRA", "RUS", "MNG", "MEX", "CAN", "PAN", "ECU", "AUS", "SRB", "BWA"} <= set(actual_codes)

    projects = workbook["projects"]
    country_input_column = [
        cell.value for cell in projects[1]
    ].index("country_iso3") + 1
    country_letter = projects.cell(1, country_input_column).column_letter
    validation = next(
        item
        for item in projects.data_validations.dataValidation
        if f"{country_letter}2:{country_letter}500" in str(item.sqref)
    )
    lookup_letter = lookups.cell(1, country_column).column_letter
    assert validation.formula1 == (
        f"'_LOOKUPS'!${lookup_letter}$2:${lookup_letter}${len(expected_codes) + 1}"
    )
