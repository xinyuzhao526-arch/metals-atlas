from io import BytesIO

from openpyxl import load_workbook

from app.demo import build_demo_workbook
from app.excel import build_template
from app.models import GuidanceObservation, ProductionObservation, ReserveObservation, Source


def upload_bytes(client, headers, payload: bytes, name: str = "fixture.xlsx"):
    return client.post(
        "/api/v1/admin/imports",
        files={"file": (name, BytesIO(payload), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )


def confirm_demo(client, headers):
    uploaded = upload_bytes(client, headers, build_demo_workbook(), "phase1a-demo.xlsx")
    assert uploaded.status_code == 200, uploaded.text
    confirmed = client.post(f"/api/v1/admin/imports/{uploaded.json()['id']}/confirm", headers=headers)
    assert confirmed.status_code == 200, confirmed.text


def edited_demo(record_id: str, row_version: int, **changes) -> bytes:
    workbook = load_workbook(BytesIO(build_demo_workbook()))
    sheet = workbook["production"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("record_id") + 1, record_id)
    sheet.cell(2, headers.index("row_version") + 1, row_version)
    for field, value in changes.items():
        sheet.cell(2, headers.index(field) + 1, value)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def production_row(preview: dict):
    return next(row for row in preview["rows"] if row["sheet"] == "production" and row["row_number"] == 2)


def test_admin_routes_require_login(client):
    assert client.get("/api/v1/admin/excel/template").status_code == 401
    assert client.get("/api/v1/admin/imports").status_code == 401
    assert client.get("/api/v1/admin/reviews").status_code == 401
    assert upload_bytes(client, {}, build_demo_workbook()).status_code == 401


def test_update_new_uuid_and_row_version_classification(authenticated, session):
    client, headers = authenticated
    confirm_demo(client, headers)
    first = session.query(ProductionObservation).filter_by(record_key="escondida-fy2026-production").one()
    second = session.query(ProductionObservation).filter_by(record_key="morenci-2025-production").one()

    changed = edited_demo(first.id, first.row_version, original_value=124, normalized_value=124)
    update_job = upload_bytes(client, headers, changed).json()["id"]
    update_preview = client.get(f"/api/v1/admin/imports/{update_job}/preview").json()
    update_row = production_row(update_preview)
    assert update_row["classification"] == "update"
    assert update_row["raw"]["normalized_value"] == 124
    assert update_row["normalized"]["normalized_value"] == "124"
    assert update_row["before"]["normalized_value"] == "123.456000"
    assert update_row["diff"]["normalized_value"] == {
        "before": "123.456000",
        "after": "124",
    }
    assert update_preview["sheet_counts"]["production"] == 2
    assert update_preview["action_summary"]["update"] == 1
    assert update_preview["action_summary"]["pending_review_observations"] == 1

    wrong_uuid = edited_demo(second.id, first.row_version)
    uuid_job = upload_bytes(client, headers, wrong_uuid).json()["id"]
    uuid_preview = client.get(f"/api/v1/admin/imports/{uuid_job}/preview").json()
    assert production_row(uuid_preview)["classification"] == "conflict"

    stale = edited_demo(first.id, first.row_version + 1)
    stale_job = upload_bytes(client, headers, stale).json()["id"]
    stale_preview = client.get(f"/api/v1/admin/imports/{stale_job}/preview").json()
    assert production_row(stale_preview)["classification"] == "conflict"
    assert any("row_version" in error for error in production_row(stale_preview)["errors"])

    new_key = edited_demo("", first.row_version, record_key="escondida-new-stable-key")
    new_job = upload_bytes(client, headers, new_key).json()["id"]
    new_preview = client.get(f"/api/v1/admin/imports/{new_job}/preview").json()
    assert production_row(new_preview)["classification"] == "new"


def test_source_code_matches_source_relation(authenticated, session):
    client, headers = authenticated
    confirm_demo(client, headers)
    for model in (ProductionObservation, GuidanceObservation, ReserveObservation):
        for observation in session.query(model):
            source = session.get(Source, observation.source_id)
            assert source is not None
            assert observation.source_code == source.code


def test_overlapping_ownership_import_is_rejected(authenticated):
    client, headers = authenticated
    confirm_demo(client, headers)
    workbook = load_workbook(BytesIO(build_template()))
    workbook["ownership"].append([None, None, "escondida", "BHP", 50, "2025-06-01", None])
    output = BytesIO()
    workbook.save(output)
    uploaded = upload_bytes(client, headers, output.getvalue(), "overlap.xlsx")
    assert uploaded.status_code == 200, uploaded.text
    job_id = uploaded.json()["id"]
    preview = client.get(f"/api/v1/admin/imports/{job_id}/preview").json()
    assert preview["summary"] == {
        "new": 1,
        "update": 0,
        "no_change": 0,
        "conflict": 0,
        "invalid": 0,
    }
    confirmed = client.post(f"/api/v1/admin/imports/{job_id}/confirm", headers=headers)
    assert confirmed.status_code == 409
    assert "重叠" in confirmed.json()["detail"]


def test_rejected_observations_never_become_public(authenticated):
    client, headers = authenticated
    confirm_demo(client, headers)
    reviews = client.get("/api/v1/admin/reviews").json()
    for item in reviews:
        response = client.post(
            f"/api/v1/admin/reviews/{item['id']}/reject",
            headers=headers,
            json={"comment": "fixture rejected"},
        )
        assert response.status_code == 200
    assert client.get("/api/v1/public/projects").json() == []


def test_import_preview_requires_login(authenticated):
    client, headers = authenticated
    uploaded = upload_bytes(client, headers, build_demo_workbook())
    assert uploaded.status_code == 200
    client.cookies.clear()

    response = client.get(
        f"/api/v1/admin/imports/{uploaded.json()['id']}/preview"
    )

    assert response.status_code == 401
